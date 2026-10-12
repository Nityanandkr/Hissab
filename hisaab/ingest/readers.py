"""Read a statement file into a raw GRID (list of rows of cell values), without assuming where the table starts.

Real bank files have junk above the table (account name, period ...), junk below it (closing balance, disclaimers) and
come in encodings / delimiters / formats of every kind.  Finding the table is the mapper's job; this module only turns
bytes into rows.
"""
from __future__ import annotations

import csv
import io
import shutil
import subprocess
from pathlib import Path
from typing import Any

MAX_ROWS = 200_000
Grid = list[list[Any]]


class ReadError(Exception):
    """A file we cannot read; the message is safe to show to the user."""


_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
_DELIMS = (",", ";", "\t", "|")


def _decode(data: bytes) -> str:
    for enc in _ENCODINGS:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _pick_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines()[:60] if ln.strip()]
    best, best_score = ",", -1.0
    for d in _DELIMS:
        counts = [ln.count(d) for ln in lines]
        multi = [c for c in counts if c >= 2]
        if not multi:
            continue
        mode = max(set(multi), key=multi.count)
        score = multi.count(mode) + 0.01 * mode
        if score > best_score:
            best, best_score = d, score
    return best


def _read_csv(data: bytes) -> Grid:
    text = _decode(data)
    delim = _pick_delimiter(text)
    rows: Grid = []
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        if i >= MAX_ROWS:
            raise ReadError(f"File has more than {MAX_ROWS:,} rows.")
        rows.append([c.strip() for c in row])
    return rows


def _read_xlsx(data: bytes) -> list[tuple[str, Grid]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise ReadError("Reading .xlsx needs the 'openpyxl' package (pip install openpyxl).") from e
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ReadError("This .xlsx file could not be opened (is it password-protected or corrupted?).") from e
    out = []
    for ws in wb.worksheets:
        rows: Grid = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                raise ReadError(f"Sheet '{ws.title}' has more than {MAX_ROWS:,} rows.")
            rows.append(list(row))
        out.append((ws.title, rows))
    return out


def _read_xls(data: bytes) -> list[tuple[str, Grid]]:
    try:
        import xlrd
    except ImportError as e:  # pragma: no cover
        raise ReadError("Reading legacy .xls needs the 'xlrd' package (pip install xlrd).") from e
    try:
        wb = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise ReadError("This .xls file could not be opened (is it password-protected or corrupted?).") from e
    out = []
    for sh in wb.sheets():
        rows: Grid = []
        for r in range(min(sh.nrows, MAX_ROWS)):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
                    except Exception:
                        row.append(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _read_pdf(data: bytes) -> Grid:
    """Best-effort: needs poppler's `pdftotext` on the PATH. Layouts differ wildly between banks; see ingest/pdftable.py."""
    from . import pdftable
    return pdftable.pdf_to_grid(data)


def read_grids(data: bytes, filename: str) -> list[tuple[str, Grid]]:
    """Return [(sheet_name, grid)].  CSV/PDF files have one pseudo-sheet."""
    if not data:
        raise ReadError("The file is empty.")
    ext = Path(filename or "").suffix.lower()
    if data[:4] == b"PK\x03\x04" or ext in (".xlsx", ".xlsm"):
        return _read_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or ext == ".xls":
        return _read_xls(data)
    if data[:4] == b"%PDF" or ext == ".pdf":
        return [("pdf", _read_pdf(data))]
    if ext in (".csv", ".txt", ".tsv", ".tab", "") or True:
        return [("csv", _read_csv(data))]
    raise ReadError("Unsupported file type.")  # pragma: no cover
